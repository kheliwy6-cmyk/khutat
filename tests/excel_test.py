"""تصدير الخطة إلى اكسل:
- بالإعداد الأصلي (الجدول بالأيام): مثل النسخة السابقة تماماً.
- بطريقة ملف التسجيل اليومي (من لوحة التحكم): ورقة لكل خطة كما في «بيانات الخطط» (جزء/من/إلى، تجاوز الإجمالي، الصيغ)،
  وورقة الخطة بالأيام؛ والصيغ تُحسب في LibreOffice كما هي محسوبة في الملف؛ والخطة الطويلة في أوراق؛ ورفع الملف يعيد الخطة.
"""
import asyncio, io, json, re, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from harness import *
from features_test import set_start, append
from playwright.async_api import async_playwright
import openpyxl

check = Checker()
AR = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')
DAILY_SEED = {'config/settings': {'settings': {'excel': 'daily'}, 'savedAt': '2026-10-15T09:00:00.000Z'}}


def wb_of(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def rng(s):
    m = re.match(r'^(?:(\d+)/)?(\d+)(?:-(\d+))?$', s)
    return (int(m.group(1)) if m.group(1) else None, int(m.group(2)), int(m.group(3) or m.group(2))) if m else None


def expect_cells(text, spans, hint=None):
    """نص خانة (كما في الموقع) ← [جزء، من، إلى، تجاوز] كما يجب أن يكون في ملف التسجيل اليومي"""
    t = text.translate(AR).replace('‏', '').replace(' ', '')
    if not t or t == '-':
        return ['-', '-', '-', None]
    if '←' in t:
        a, b = t.split('←')
        ra, rb = rng(a), rng(b)
        sa = next(s for s in spans if s[0] == ra[0] and s[1] <= ra[1] <= s[2])
        sb = next(s for s in spans if s[0] == rb[0] and s[1] <= rb[2] <= s[2])
        if ra[0] == rb[0]:
            raise ValueError('same part')
        return [f'{ra[0]}←{rb[0]}', ra[1], rb[2], (sa[2] - ra[1] + 1) + (rb[2] - sb[1] + 1)]
    pieces = [rng(x) for x in t.split('+')]
    lv = pieces[0][0]
    if lv is None and len(pieces) > 1 and hint not in (None, '-'):
        # «١٣٣-١٤٠+٢٥-٢٦» يُكتب في الخطة بلا جزء؛ والجزء الصحيح ما فيه كل قطعة منها في متن من متونه
        inpart = lambda pg: any(sp[0] == hint and sp[1] <= pg <= sp[2] for sp in spans)
        if all(inpart(pc[1]) and inpart(pc[2]) for pc in pieces):
            lv = hint
    over = sum(p[2] - p[1] + 1 for p in pieces) if len(pieces) > 1 else None
    return ['-' if lv is None else lv, pieces[0][1], pieces[-1][2], over]


def num(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


async def export(page, out):
    name, data = await click_save(page, '#exportXlsx', OUT / out)
    return name, data


async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        # ١) الإعداد الأصلي مثل النسخة السابقة
        same = []
        for build in ('old', 'new'):
            ctx, page, errs = await open_page(b, build, owner=True)
            await choose_level(page, 2)
            _, data = await export(page, f'simple-{build}.xlsx')
            wb = wb_of(data)
            same.append((wb.sheetnames, values(wb.worksheets[0])))
            await ctx.close()
        check(same[0] == same[1], f'default Excel export unchanged: sheets {same[1][0]}, {len(same[1][1])} rows')

        # ٢) بطريقة ملف التسجيل اليومي
        ctx, page, errs = await open_page(b, owner=True, seed=DAILY_SEED)
        spans = await page.evaluate("""TEXTS.map(t => { const us = PlanCore.buildUnits(t.days); return [us[0].level, us[0].page, us[us.length - 1].page]; })""")
        await choose_level(page, 2)
        days_rows = (await dump(page))['rows']
        name, data = await export(page, 'daily-level2.xlsx')
        wb = wb_of(data)
        check(wb.sheetnames == ['للصق في بيانات الخطط', 'الخطة بالأيام'], f'daily format: sheets {wb.sheetnames}')
        ws = wb.worksheets[0]
        head = [ws['A1'].value, ws['D1'].value, ws['A2'].value, ws['D2'].value, ws['A3'].value, ws['G1'].value]
        check(head == ['اسم/رقم المستوى:', 'المستوى 2', 'اسم المتن:', 'الأربعون النووية-نواقض الإسلام-القواعد الأربع-الأصول الثلاثة', 'عدد مقادير الخطة (تلقائي):', 'إجمالي صفحات الحفظ:'],
              f'block head as in «بيانات الخطط»: {head}')
        sub = [ws.cell(5, c).value for c in range(2, 17)]
        check(sub == ['جزء', 'من', 'إلى', 'جزء', 'من', 'إلى', 'جزء', 'من', 'إلى', 'حفظ', 'ربط', 'كبرى', 'حفظ', 'ربط', 'كبرى'], 'column headers جزء/من/إلى and the totals')
        merges = sorted(str(m) for m in ws.merged_cells.ranges)
        check(all(m in merges for m in ['A1:C1', 'D1:F1', 'A4:A5', 'B4:D4', 'N4:P4']), 'merged cells as in the template')
        bad = []
        for i, r in enumerate(days_rows):
            row = 6 + i
            got = [num(ws.cell(row, c).value) for c in range(2, 14)]
            want = []
            for j, text in enumerate(r[:3]):
                want.append(expect_cells(text, spans, got[3 * j]))
            exp = [w for x in want for w in x[:3]] + [x[3] for x in want]
            if got != exp or ws.cell(row, 1).value != i + 1:
                bad.append((i + 1, r[:3], got, exp))
        check(not bad, f'each day → جزء/من/إلى and تجاوز الإجمالي ({len(days_rows)} days) {bad[:2]}')
        cross = [i for i, r in enumerate(days_rows) if '←' in r[1] or '+' in r[1]]
        check(len(cross) > 0, f'the plan has linking that spans parts ({len(cross)} days), written with تجاوز الإجمالي')
        f6 = ws['N6'].value
        check(f6 == '=IFERROR(IF(K6<>"",K6,IF(C6="","",D6-C6+1)),"")' and ws['I1'].value == '=IFERROR(SUM(N6:N105),0)' and ws['D3'].value == '=COUNT(N6:N105)',
              f'the same formulas: {f6}')
        check(ws['A105'].value == 100 and ws['B48'].value is None, 'rows 6–105 numbered 1–100, empty after the last day')
        # الصيغ في LibreOffice
        import lo
        L = lo.LO()
        try:
            doc = L.open(OUT / 'daily-level2.xlsx')
            doc.calculateAll()
            sh = doc.Sheets.getByIndex(0)
            badf = []
            for i, r in enumerate(days_rows):
                row = 6 + i
                for j, (a, bb, o, n) in enumerate([('C', 'D', 'K', 'N'), ('F', 'G', 'L', 'O'), ('I', 'J', 'M', 'P')]):
                    ov = ws[f'{o}{row}'].value
                    fa, fb = ws[f'{a}{row}'].value, ws[f'{bb}{row}'].value
                    want = ov if ov is not None else (fb - fa + 1 if isinstance(fa, int) and isinstance(fb, int) else '')
                    got = num(lo.value(sh, f'{n}{row}'))
                    if (got if got != '' else None) != (want if want != '' else None) or (isinstance(got, (int, float)) and got <= 0):
                        badf.append((row, n, got, want))
            tot = [num(lo.value(sh, x)) for x in ('I1', 'I2', 'I3', 'D3')]
            check(not badf, f'LibreOffice computes الإجمالي الفعلي for every day, all positive {badf[:3]}')
            memo_pages = sum(int(re.split('[-]', x)[-1]) - int(re.split('[-/]', x)[-2]) + 1 for x in [])
            check(tot[3] == len(days_rows) and all(isinstance(x, (int, float)) and x > 0 for x in tot), f'totals and count computed: {tot}')
        finally:
            L.close()
        cached = [ws['I1'].value]

        # ٣) خطة أطول من ١٠٠ يوم: ورقتان
        await choose_level(page, 5)
        await append(page, 'bulugh')
        n_days = len((await dump(page))['rows'])
        _, data = await export(page, 'daily-long.xlsx')
        wb = wb_of(data)
        s1, s2 = wb.worksheets[0], wb.worksheets[1]
        check(wb.sheetnames == ['للصق 1 من 2', 'للصق 2 من 2', 'الخطة بالأيام'] and s1['D2'].value.endswith('(1 من 2)') and s2['D2'].value.endswith('(2 من 2)'),
              f'{n_days} days → two plan sheets: {wb.sheetnames}, {s2["D2"].value}')
        n2 = sum(1 for r in range(6, 106) if s2.cell(r, 2).value is not None)
        check(n2 == n_days - 100 and s1.cell(105, 2).value is not None, f'second sheet holds the remaining {n2} days')
        check(s1['R1'].value and 'طريقة إضافة' in s1['R1'].value and s1['T1'].value == 'رقم الخطة' and s1['U11'].value == 1 + 106 * 9 and s2['R1'].value is None,
              'paste instructions and the plan → first-row table on the first sheet only')

        # ٤) البدء من صفحة، وخانة معدَّلة باليد
        await choose_level(page, 5)
        await set_checked(page, ['tahawiya'])
        await set_start(page, 'tahawiya', 'from', 150)
        await page.fill('#c-0-memo', '٥/١٥٠-١٥١')
        await page.dispatch_event('#c-0-memo', 'input')
        await page.dispatch_event('#c-0-memo', 'change')
        await page.wait_for_timeout(200)
        _, data = await export(page, 'daily-tahawiya.xlsx')
        ws = wb_of(data).worksheets[0]
        check(ws['D2'].value == 'العقيدة الطحاوية من ص150' and [ws['B6'].value, ws['C6'].value, ws['D6'].value] == [5, 150, 151],
              f'start page in the name, and an edited cell exported as edited: {ws["D2"].value}, {[ws["B6"].value, ws["C6"].value, ws["D6"].value]}')

        # ٥) رفع الملف يعيد الخطة (من ورقة الخطة بالأيام)، ورفع أوراق «بيانات الخطط» وحدها
        await choose_level(page, 3)
        await set_checked(page, ['tawheed'])
        orig = (await dump(page))['rows']
        _, data = await export(page, 'daily-tawheed.xlsx')
        await choose_level(page, 2)
        await page.set_input_files('#importFile', str(OUT / 'daily-tawheed.xlsx'))
        await page.wait_for_timeout(600)
        back = (await dump(page))['rows']
        check(back == orig, f'uploading the exported file brings the plan back ({len(back)} days)')
        wb = wb_of(data)
        del wb['الخطة بالأيام']
        wb.save(OUT / 'daily-only.xlsx')
        await page.set_input_files('#importFile', str(OUT / 'daily-only.xlsx'))
        await page.wait_for_timeout(600)
        st = await page.inner_text('#status')
        back2 = (await dump(page))['rows']
        check([r[0] for r in back2] == [r[0] for r in orig] and 'رُفعت' in st, f'a file with only the «بيانات الخطط» sheet is read too: {st}')
        check(not errs, 'no page errors: ' + '; '.join(errs[:3]))
        await b.close()
    print('FAILURES:', len(check.fails), 'of', check.n)


if __name__ == '__main__':
    asyncio.run(main())
